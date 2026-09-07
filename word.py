# word.py 修复版本 - 主要修复表格创建部分
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import datetime

def create_complete_sentinel_document():
    """创建完整的Sentinel CPNA设计文档"""
    try:
        # 创建新文档
        doc = Document()
        
        # 设置文档属性
        doc.core_properties.title = "Sentinel CPNA 自动化系统设计文档"
        doc.core_properties.author = "网络自动化团队"
        doc.core_properties.subject = "物理感知网络自动化系统技术规整"
        
        # 1. 封面页
        add_cover_page(doc)
        
        # 2. 目录
        add_table_of_contents(doc)
        
        # 3. 各章节内容
        add_chapter_1(doc)  # 初始需求说明书
        add_chapter_2(doc)  # 理想功能需求清单
        add_chapter_3(doc)  # 需求完成度对照表
        add_chapter_4(doc)  # 系统当前功能说明
        add_chapter_5(doc)  # 未来系统架构与技术计划
        
        # 保存文档
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"Sentinel_CPNA_Design_Document_{timestamp}.docx"
        doc.save(filename)
        
        print(f"✅ 文档创建成功: {filename}")
        print(f"📁 文件路径: {os.path.abspath(filename)}")
        
        return filename
        
    except Exception as e:
        print(f"❌ 创建文档时出错: {e}")
        import traceback
        print(f"详细错误: {traceback.format_exc()}")
        return None

def add_cover_page(doc):
    """添加封面页"""
    # 添加标题
    title = doc.add_heading('Sentinel CPNA 自动化系统', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 添加副标题
    subtitle = doc.add_heading('物理感知、分层驱动的网络自动化系统', 1)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 添加版本信息
    doc.add_paragraph()
    version_info = doc.add_paragraph()
    version_info.add_run('版本: ').bold = True
    version_info.add_run('v2.0 (完整技术规整版)')
    
    # 添加日期
    date_info = doc.add_paragraph()
    date_info.add_run('创建日期: ').bold = True
    date_info.add_run(datetime.datetime.now().strftime("%Y年%m月%d日"))
    
    # 添加分隔线
    doc.add_paragraph("_" * 80)
    
    # 添加项目愿景
    vision = doc.add_paragraph()
    vision.add_run('项目愿景:').bold = True
    doc.add_paragraph('构建一个物理感知的、分层驱动的网络自动化系统。借鉴 CAE 行业（如 ANSYS, Abaqus）对复杂仿真任务的分段处理逻辑，通过简单的标记语言（Sentinel Markup）管理跨厂商设备的复杂网络配置，并具备感知物理环境（温度、负载）并自动调整执行节奏的能力。')
    
    # 分页
    doc.add_page_break()

def add_table_of_contents(doc):
    """添加目录"""
    heading = doc.add_heading('目录', 1)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    
    toc_items = [
        "1. 初始需求说明书",
        "   1.1 项目愿景",
        "   1.2 核心功能需求",
        "2. 理想功能需求清单",
        "3. 需求完成度对照表",
        "4. 系统当前功能说明清单",
        "5. 未来系统架构与技术计划",
        "   5.1 推荐的项目目录结构",
        "   5.2 核心技术讨论要点与迁移路线",
        "   5.3 主要功能与重点函数详细说明",
        "   5.4 当前需要深入讨论的迁移遗留技术讨论要点",
        "6. 总结：Sentinel 终极三层架构"
    ]
    
    for item in toc_items:
        doc.add_paragraph(item)
    
    doc.add_page_break()

def add_chapter_1(doc):
    """添加第1章：初始需求说明书"""
    doc.add_heading('1. 初始需求说明书', 1)
    
    # 1.1 项目愿景
    doc.add_heading('1.1 项目愿景', 2)
    doc.add_paragraph('构建一个物理感知的、分层驱动的网络自动化系统。借鉴 CAE 行业（如 ANSYS, Abaqus）对复杂仿真任务的分段处理逻辑，通过简单的标记语言（Sentinel Markup）管理跨厂商设备的复杂网络配置，并具备感知物理环境（温度、负载）并自动调整执行节奏的能力。')
    
    # 1.2 核心功能需求
    doc.add_heading('1.2 核心功能需求', 2)
    
    # F01: 分层标记解析
    doc.add_heading('F01: 分层标记解析 (Hierarchical Parsing)', 3)
    p = doc.add_paragraph()
    p.add_run('支持四级逻辑隔离标记：').bold = True
    
    # 使用项目符号列表
    doc.add_paragraph('• # (Set)：命令集结束符（如接口、路由配置等）', style='List Bullet')
    doc.add_paragraph('• ## (Category)：命令类别结束符（如 config, show, debug 等）', style='List Bullet')
    doc.add_paragraph('• ### (Device)：单个设备配置结束符', style='List Bullet')
    doc.add_paragraph('• #### (File)：整个任务文件结束符', style='List Bullet')
    
    doc.add_paragraph('动态上下文切换：解析器需在读取设备元数据（IP/Port）与多层级命令流之间无缝切换。')
    
    # F02: 物理场语义支持
    doc.add_heading('F02: 物理场语义支持 (Physics-Aware Enums)', 3)
    p = doc.add_paragraph()
    p.add_run('设备类型枚举：').bold = True
    p.add_run('识别 CISCO_ROUTER, CISCO_SWITCH, HUAWEI_ROUTER, FIREWALL, WINDOWS_PC, LINUX_SERVER, GENERIC 等。')
    
    p = doc.add_paragraph()
    p.add_run('设备状态枚举：').bold = True
    p.add_run('DISCONNECTED → CONNECTED → CONFIGURED → ERROR。')
    
    p = doc.add_paragraph()
    p.add_run('配置类别枚举：').bold = True
    p.add_run('INTERFACE_CONFIG, ROUTING_CONFIG, SECURITY_CONFIG, SYSTEM_CONFIG, DIAGNOSTIC, MONITORING。')
    
    # F03: 模板化与变量执行
    doc.add_heading('F03: 模板化与变量执行 (Template Execution)', 3)
    doc.add_paragraph('集成 Jinja2 引擎，支持在配置文件中通过 {{ variable }} 进行变量替换，满足大规模差异化部署。')
    
    # F04: 项目级环境隔离
    doc.add_heading('F04: 项目级环境隔离 (Project Isolation)', 3)
    doc.add_paragraph('强制闭环管理：所有 Snapshot、Log、Output 必须严格限定在当前运行文件所属的项目目录下，禁止文件散落在根目录。')

def add_chapter_2(doc):
    """添加第2章：理想功能需求清单"""
    doc.add_heading('2. 理想功能需求清单', 1)
    doc.add_paragraph('针对未来网络孪生与工业级稳定性的需求规划。')
    
    # 创建表格 - 修复这里的问题
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # 设置表头
    header_cells = table.rows[0].cells
    headers = ['需求 ID', '功能分类', '需求描述', '核心价值']
    
    for i, header in enumerate(headers):
        header_cells[i].text = header
        header_cells[i].paragraphs[0].runs[0].font.bold = True
    
    # 添加数据行 - 确保行数足够
    data_rows = [
        ['REQ-01', '资产管理', '支持通过结构化 YAML 或自定义 # 标记文件批量导入设备资产。', '灵活的多模输入'],
        ['REQ-02', '并发控制', '基于 ThreadPoolExecutor 并向 Asyncio 演进，具备超时处理机制，防止单台卡死全局。', '高效率执行'],
        ['REQ-03', '项目管理', '运行结果、备份、日志自动归档至 [Project_Dir]/outputs/ 下，互不干扰。', '目录结构清晰'],
        ['REQ-04', '容错与安全', '正式变更前执行 initial_config 备份；失败时触发回滚。', '生产环境保命符'],
        ['REQ-05', '多厂商适配', '依据设备类型自动切换 SSH/Telnet，并适配厂商特定的 CLI 差异。', '跨平台兼容性'],
        ['REQ-06', '审计与报告', '生成人类可读 (TXT) 和机器可读 (JSON) 的双重报告。', '部署结果可回溯'],
        ['IR-01', '语义化分段器', '解析器识别命令边界，将其转化为带延时标签的"执行块"。', '解决执行节奏问题'],
        ['IR-02', '混合输入引擎', '程序预处理阶段将 YAML 列表映射为带 # 逻辑的原子操作。', '提升书写效率'],
        ['IR-03', '物理场快照管理', '每次变更前在 Project/Snapshots/ 下生成版本化的 .conf。', '物理存证'],
        ['IR-04', '智能回滚判定', '支持 Post-check 逻辑（如：Ping 判定、路由表状态判定触发回滚）。', '提升交付质量']
    ]
    
    for row_data in data_rows:
        # 添加新行
        row = table.add_row()
        # 填充单元格数据
        for i, cell_data in enumerate(row_data):
            # 确保单元格存在
            if i < len(row.cells):
                row.cells[i].text = cell_data
            else:
                # 如果单元格不够，先添加
                while len(row.cells) <= i:
                    row.add_cell()
                row.cells[i].text = cell_data

def add_chapter_3(doc):
    """添加第3章：需求完成度对照表"""
    doc.add_heading('3. 需求完成度对照表', 1)
    doc.add_paragraph('功能模块实现状态与原始构想的对比分析。')
    
    # 创建表格
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    
    # 设置表头
    header_cells = table.rows[0].cells
    headers = ['功能模块', '原始构想 (# 标记逻辑)', '现状 (v1.x 代码实现)', '状态', '差距分析 (Gap Analysis)']
    
    for i, header in enumerate(headers):
        header_cells[i].text = header
        header_cells[i].paragraphs[0].runs[0].font.bold = True
    
    # 添加数据行
    data_rows = [
        ['配置解析', '四级标记 (# 系列)', 'YAML ConfigParser', '⚡ 30%', '尚未实现对 # 文本流的解析，目前仅依赖 YAML 列表。'],
        ['并发执行', '异步协程驱动', '多线程 ThreadPool', '✅ 70%', '已有并发雏形，但尚未实现基于"块"的延时注入。'],
        ['项目管理', '目录完全隔离', '动态路径识别', '✅ 90%', '已实现项目路径跟随，仅差 Snapshot 路径强制归口。'],
        ['容错回滚', '主动逻辑判定回滚', '异常触发被动回滚', '⚡ 50%', '缺乏"验证不通过则回滚"的深度判定机制。'],
        ['逆向工程', '自动提取并生成标记文本', '备份转 YAML', '⚡ 60%', '清洗逻辑需加强，需能自动过滤无用指令。']
    ]
    
    for row_data in data_rows:
        row = table.add_row()
        for i, cell_data in enumerate(row_data):
            if i < len(row.cells):
                row.cells[i].text = cell_data
            else:
                while len(row.cells) <= i:
                    row.add_cell()
                row.cells[i].text = cell_data

def add_chapter_4(doc):
    """添加第4章：系统当前功能说明清单"""
    doc.add_heading('4. 系统当前功能说明清单', 1)
    doc.add_paragraph('基于目前 main.py 及相关 src/ 代码的实际功能总结：')
    
    # 使用编号列表
    functions = [
        '项目级环境隔离运行：系统通过命令行参数锁定工作目录，确保每次实验的结果在物理路径上闭环。',
        '对象化资产建模：利用 Device 和 Command 模型，将复杂的网络参数结构化存储。',
        '多协议自适应：支持 Cisco/Huawei 的 Telnet 和 SSH 动态映射与连接管理。',
        '非侵入式前置快照：在部署开始前自动抓取 running-config 并生成带时间戳的 .conf 备份。',
        '异常触发式自动恢复：内置故障回滚逻辑，在连接失败或执行异常时尝试还原初始状态。',
        '结构化部署审计：输出包含详细回显的 TXT 报告和用于系统集成的 JSON 汇总表。'
    ]
    
    for i, func in enumerate(functions, 1):
        doc.add_paragraph(f'{i}. {func}')

def add_chapter_5(doc):
    """添加第5章：未来系统架构与技术计划"""
    doc.add_heading('5. 未来系统架构与技术计划', 1)
    
    # 5.1 推荐的项目目录结构
    doc.add_heading('5.1 推荐的项目目录结构 (Sentinel CPNA)', 2)
    
    # 使用等宽字体显示目录结构
    dir_structure = """network_automation/
├── configs/                # 原始输入目录 (devices.txt / devices.yaml)
├── logs/                   # 全局与项目日志
├── outputs/                # 结果输出
│   └── [project_name]/     # 按项目隔离
│       ├── snapshots/      # 存放变更前的 .conf 文件
│       └── reports/        # JSON/TXT 报告
├── src/
│   ├── core/               # 核心层
│   │   ├── config_parser.py        # 支持 # 标记的状态机解析器
│   │   ├── dynamic_config_loader.py # 动态模板加载
│   │   ├── device_manager.py       # 连接与备份管理
│   │   ├── command_executor.py     # 块状指令执行引擎 (实现延时注入)
│   │   └── physics_engine.py       # 物理感知与 ODE 反馈算法
│   ├── models/             # 模型层 (Device, Command, Enums)
│   └── utils/              # 工具层 (Logger, Validators)
├── main.py                 # 入口，参数解析与调度
└── requirements.txt        # 依赖包"""
    
    code_para = doc.add_paragraph(dir_structure)
    code_para.style = 'Normal'
    for run in code_para.runs:
        run.font.name = 'Consolas'
        run.font.size = Pt(10)
    
    # 5.2 核心技术讨论要点与迁移路线
    doc.add_heading('5.2 核心技术讨论要点与迁移路线 (Roadmap)', 2)
    
    topics = [
        {
            'title': '分层标记解析器的实现',
            'problem': '如何替代简单的 YAML？',
            'plan': '开发一个类似 CAE 输入解析的状态机，读取文本流，遇到 # 即封装为一个 CommandSet，遇到 ## 切换模式。'
        },
        {
            'title': '语义化延时注入 (Execution Cadence)',
            'problem': '网络协议收敛的物理延时问题',
            'plan': '在 YAML 预处理或 # 解析时，识别 interface、router ospf 等关键字，自动在执行队列中注入 sleep 动作。'
        },
        {
            'title': '智能回滚准则 (Active Rollback)',
            'problem': '仅依赖报错回滚不够智能',
            'plan': '引入"后置验证"块，若 show ip int brief 结果不符合预期，即使没有 SSH 报错也要强制回滚。'
        },
        {
            'title': 'Snapshot 路径强制规范',
            'problem': '备份文件散落在根目录',
            'plan': '修改 DeviceManager 的存储逻辑，强制将备份文件从根目录移入 outputs/snapshots/ 目录下。'
        }
    ]
    
    for topic in topics:
        doc.add_heading(topic['title'], 3)
        p = doc.add_paragraph()
        p.add_run('问题：').bold = True
        p.add_run(f' {topic["problem"]}')
        
        p = doc.add_paragraph()
        p.add_run('计划：').bold = True
        p.add_run(f' {topic["plan"]}')
    
    # 5.3 主要功能与重点函数详细说明
    doc.add_heading('5.3 主要功能与重点函数详细说明', 2)
    doc.add_paragraph('本章节是系统的"技术护城河"。在论文或面试中，需重点强调这三个模块，它们直接体现了从普通自动化脚本向"工业仿真级"系统的跨越：')
    
    modules = [
        {
            'name': 'A. 物理反馈模块 (physics_engine.py)',
            'desc': '这是 Sentinel 的核心护城河，负责将瞬息万变的物理环境指标转化为确定性的执行策略。',
            'functions': [
                {
                    'name': 'predict_congestion_risk(temp, cpu_load)',
                    'desc': '输入实时采集的设备温度和 CPU 负载，利用预设的 ODE（常微分方程）降维模型，计算出设备在未来 N 秒内发生 I/O 阻塞或响应超时的概率。'
                },
                {
                    'name': 'get_dynamic_window_size(health_score)',
                    'desc': '根据物理健康度评分，实时返回一个动态并发窗口数。示例：健康度为 100（环境优良）时，并发数设为 50；当感知到局部热点导致健康度降至 40 时，自动缩减并发至 5，确保不压垮设备。'
                }
            ]
        },
        {
            'name': 'B. 异步调度模块 (dispatcher.py)',
            'desc': '体现处理大规模并发任务与实时反馈闭环的能力。',
            'functions': [
                {
                    'name': 'async execute_batch_tasks(task_list)',
                    'desc': '利用 Python 的 asyncio.gather 或 TaskGroup 启动非阻塞任务流，实现指令的高并发植入。'
                },
                {
                    'name': 'async monitor_and_adjust()',
                    'desc': '后台常驻协程。每秒监听 physics_engine 的输出，若发现风险指标升高，立即动态修改 dispatcher 的下发速率。这是系统的"中枢神经"。'
                }
            ]
        },
        {
            'name': 'C. 厂商适配模块 (adapters/)',
            'desc': '体现对 Model-Driven Networking（模型驱动网络） 的理解与应用。',
            'functions': [
                {
                    'name': 'translate_to_yang(intent_json)',
                    'desc': '将高层业务意图（如"隔离受攻击端口"）翻译为厂商中立的 YANG 数据结构。'
                },
                {
                    'name': 'async set_config_gnmi(payload)',
                    'desc': '利用 pygnmi 库，将 JSON 化的配置通过 gRPC 协议推送到设备，彻底告别脆弱的 CLI 交互。'
                }
            ]
        }
    ]
    
    for module in modules:
        doc.add_heading(module['name'], 3)
        doc.add_paragraph(module['desc'])
        
        for func in module['functions']:
            p = doc.add_paragraph()
            p.add_run('• 重点函数：').bold = True
            p.add_run(f'{func["name"]}')
            
            p = doc.add_paragraph()
            p.add_run('  功能：').bold = True
            p.add_run(f'{func["desc"]}')
    
    # 5.4 当前需要深入讨论的迁移遗留技术讨论要点
    doc.add_heading('5.4 当前需要深入讨论的迁移遗留技术讨论要点', 2)
    doc.add_paragraph('为了实现"伟大系统"的目标，重构必须紧扣："物理感知、异步驱动、中立抽象"。')
    
    discussion_points = [
        {
            'title': '1. 内核重构：异步 I/O 引擎',
            'points': [
                '核心方案：全面从多线程（Thread）转向纯 asyncio 协程驱动。',
                '技术替代：利用 gNMI 和 Netconf 替换老旧的 CLI 交互。',
                '终极目标：实现百万级指令的非阻塞下发，彻底解决 PNETLab 实验环境或大规模生产环境下的 I/O 瓶颈。'
            ]
        },
        {
            'title': '2. 插件式感知层：物理-逻辑耦合 (The Physics Plugin)',
            'points': [
                '数学模型：引入 ODE 降维模型（借鉴 Modelica 思路），不再追求复杂的全量 PDE 计算，而是快速建立"温度-电力-负载"与"网络延迟"的映射。',
                '虚拟传感器 (Mock Sensor)：在代码中预留接口，支持模拟环境数据输入。这使得 Sentinel 具备"感知物理场"并"动态收缩并发窗口"的自愈本能。'
            ]
        },
        {
            'title': '3. 资产与背景建模 (CMDB Integration)',
            'points': [
                '上下文注入：集成 NetBox 等 CMDB 系统。',
                '约束逻辑：引入设备寿命、物理位置（如是否在高温机架）、维保记录作为决策约束，为自动化操作注入"历史纵深感"。'
            ]
        },
        {
            'title': '4. 深度解析：三位一体的感知路径',
            'subpoints': [
                {
                    'title': 'A. 状态采集：从"推"到"订阅" (Streaming Telemetry)',
                    'points': [
                        'gNMI：推荐路径。基于 HTTP/2 支持推送模式，订阅 CPU/温度路径，设备异动即触发 Sentinel 响应。',
                        'SNMP MIBs (备选)：针对老旧设备，读取 entitySensorMIB 捕获物理指标。',
                        'On-box Scripts：支持设备本地计算平均值后再上报，降低网络带宽压力。'
                    ]
                },
                {
                    'title': 'B. CMDB 集成：赋予设备"生命背景"',
                    'points': [
                        '静态数据引入：评估健康度时，计入服役年限（老化因子）和维保记录。',
                        '环境基准：识别设备是处于标准冷池还是恶劣工厂车间，调整其风险阈值。'
                    ]
                },
                {
                    'title': 'C. 外部感知器：物联网化 (MQTT Integration)',
                    'points': [
                        '物理层捕获：通过 MQTT 协议订阅机柜冷热通道传感器、震动传感器的实时数据。',
                        '交叉验证 (Cross-Verification)：若外部传感器检测到环境剧变（如震动），即便设备内部传感器尚未告警，系统也应提前执行保护策略。'
                    ]
                }
            ]
        }
    ]
    
    for point in discussion_points:
        if 'subpoints' in point:
            doc.add_heading(point['title'], 3)
            for subpoint in point['subpoints']:
                doc.add_heading(subpoint['title'], 4)
                for p in subpoint['points']:
                    doc.add_paragraph(f'• {p}', style='List Bullet')
        else:
            doc.add_heading(point['title'], 3)
            for p in point['points']:
                doc.add_paragraph(f'• {p}', style='List Bullet')
    
    # 总结：Sentinel 终极三层架构
    doc.add_heading('💡 总结：Sentinel 终极三层架构', 2)
    
    layers = [
        {
            'name': '1. 感知层 (Sensory)：物理孪生节点',
            'desc': '通过 Telemetry 订阅物理参数，利用 ODE 方程推演"网络健康余量"。'
        },
        {
            'name': '2. 逻辑决策层 (Cognition)：系统的"大脑"',
            'desc': '计算最优并发数 N_opt，并处理意图翻译与冲突检测。'
        },
        {
            'name': '3. 执行反馈层 (Execution)：异步驱动中心',
            'desc': '执行配置并立即通过闭环验证（Verification Loop）判定效果，必要时触发 Atomic Rollback。'
        }
    ]
    
    for layer in layers:
        p = doc.add_paragraph()
        p.add_run(layer['name']).bold = True
        doc.add_paragraph(layer['desc'])
    
    doc.add_heading('伟大之处：', 3)
    great_points = [
        '确定性 (Determinism)：通过物理模型掌控全局。',
        '鲁棒性 (Robustness)：在 AI 进入逻辑陷阱前，感知物理异动提前避险。',
        '跨代兼容：抹平 CLI 与 YANG 的鸿沟，做万能的"翻译官"。'
    ]
    
    for point in great_points:
        doc.add_paragraph(f'• {point}', style='List Bullet')

# 主程序入口
if __name__ == "__main__":
    print("\n=== 创建完整的Sentinel CPNA设计文档 ===")
    result = create_complete_sentinel_document()
    
    if result:
        print(f"\n🎉 文档创建完成！")
        print(f"📄 文件名: {result}")
        print(f"📁 完整路径: {os.path.abspath(result)}")
        print("\n文档包含以下章节：")
        print("1. 初始需求说明书")
        print("2. 理想功能需求清单")
        print("3. 需求完成度对照表")
        print("4. 系统当前功能说明清单")
        print("5. 未来系统架构与技术计划")
    else:
        print("❌ 文档创建失败，请检查错误信息。")
